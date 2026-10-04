"""Equation conversion and portable session storage; independent of the GUI/OCR."""
from __future__ import annotations
import base64
import io
import json
import os
import re
import tempfile
from pathlib import Path
from PIL import Image, ImageOps

MATH_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'


def clean_latex(value: str) -> str:
    value = value.strip()
    for left, right in [('$$', '$$'), ('\\[', '\\]'), ('\\(', '\\)'), ('$', '$')]:
        if value.startswith(left) and value.endswith(right):
            value = value[len(left):-len(right)].strip()
            break
    if not value:
        raise ValueError('Công thức trống.')
    return value


def to_mathml(latex: str) -> str:
    from latex2mathml.converter import convert
    from lxml import etree
    value = convert(clean_latex(latex))
    root = etree.fromstring(value.encode(), etree.XMLParser(resolve_entities=False, no_network=True))
    # Unknown commands must not silently become plain text in exported equations.
    for el in root.iter():
        if el.text and re.search(r'\\[A-Za-z]+', el.text):
            raise ValueError('Lệnh LaTeX chưa hỗ trợ: ' + el.text)
    return value


def to_omml(latex: str):
    import mathml2omml
    from lxml import etree
    xml = mathml2omml.convert(to_mathml(latex))
    # Upstream 0.0.2 closes groupChrPr as groupChr for stretchy accents.
    # Repair only this known malformed property block, never parse with recovery.
    xml = re.sub(r'(<m:groupChrPr>.*?)(</m:groupChr>)',
                 r'\1</m:groupChrPr>', xml, flags=re.S)
    # mathml2omml emits m: prefixes without declaring their namespace.
    root = etree.fromstring(
        f'<root xmlns:m="{MATH_NS}">{xml}</root>'.encode(),
        etree.XMLParser(resolve_entities=False, no_network=True),
    )
    node = root[0]
    # Treat common top accents as Word accents rather than grouping glyphs.
    accents = {'~': '\u0303', '^': '\u0302', '¯': '\u0305', '˙': '\u0307', '¨': '\u0308'}
    for group in list(node.iter(f'{{{MATH_NS}}}groupChr')):
        properties = group.find(f'{{{MATH_NS}}}groupChrPr')
        if properties is None:
            continue
        char = properties.find(f'{{{MATH_NS}}}chr')
        pos = properties.find(f'{{{MATH_NS}}}pos')
        value = char.get(f'{{{MATH_NS}}}val') if char is not None else None
        if value in accents and pos is not None and pos.get(f'{{{MATH_NS}}}val') == 'top':
            group.tag = f'{{{MATH_NS}}}acc'
            properties.tag = f'{{{MATH_NS}}}accPr'
            char.set(f'{{{MATH_NS}}}val', accents[value])
            properties.remove(pos)
    if etree.QName(node).localname not in ('oMath', 'oMathPara'):
        raise ValueError('Không tạo được Equation OMML.')
    return node


def export_docx(path: str, formulas: list[dict]):
    from docx import Document
    from docx.shared import Pt
    if not formulas:
        raise ValueError('Hãy tick ít nhất một công thức.')
    # Validate all conversions before writing; no partial/quietly missing equations.
    nodes = []
    for i, formula in enumerate(formulas, 1):
        try:
            nodes.append(to_omml(formula['latex']))
        except Exception as exc:
            raise ValueError(f'Công thức {i}: {exc}') from exc
    doc = Document()
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)
    doc.add_heading('Equation Studio', 0)
    for i, (formula, node) in enumerate(zip(formulas, nodes), 1):
        doc.add_paragraph(f"{i}. {formula.get('source_name', '')}")
        doc.add_paragraph()._p.append(node)
    destination = Path(path)
    fd, temporary = tempfile.mkstemp(suffix='.docx', dir=destination.parent)
    os.close(fd)
    try:
        doc.save(temporary)
        os.replace(temporary, destination)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def image_png(path: str) -> bytes:
    with Image.open(path) as original:
        return pil_png(original)


def pil_png(original: Image.Image) -> bytes:
    image = ImageOps.exif_transpose(original).convert('RGBA')
    background = Image.new('RGBA', image.size, 'white')
    image = Image.alpha_composite(background, image).convert('RGB')
    output = io.BytesIO()
    image.save(output, format='PNG')
    return output.getvalue()


def save_session(path: str, images: list[dict], formulas: list[dict]):
    payload = {'version': 1,
               'images': [{'id': x['id'], 'name': x['name'],
                           'png': base64.b64encode(x['png']).decode()} for x in images],
               'formulas': formulas}
    destination = Path(path)
    fd, temporary = tempfile.mkstemp(dir=destination.parent, suffix='.tmp')
    try:
        with os.fdopen(fd, 'w', encoding='utf-8') as stream:
            json.dump(payload, stream, ensure_ascii=False)
        os.replace(temporary, destination)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def load_session(path: str):
    payload = json.loads(Path(path).read_text(encoding='utf-8'))
    if payload.get('version') != 1:
        raise ValueError('Phiên bản project không hỗ trợ.')
    images, ids = [], set()
    for x in payload['images']:
        if not isinstance(x['id'], str) or x['id'] in ids:
            raise ValueError('ID ảnh không hợp lệ.')
        data = base64.b64decode(x['png'], validate=True)
        with Image.open(io.BytesIO(data)) as im:
            im.verify()
        images.append({'id': x['id'], 'name': str(x['name']), 'png': data})
        ids.add(x['id'])
    formulas = payload['formulas']
    if not isinstance(formulas, list):
        raise ValueError('Danh sách công thức không hợp lệ.')
    for f in formulas:
        if f['image_id'] not in ids or not isinstance(f['latex'], str):
            raise ValueError('Công thức tham chiếu ảnh không hợp lệ.')
        box = f['box']
        if len(box) != 4 or any(not isinstance(n, (int, float)) for n in box):
            raise ValueError('Vùng công thức không hợp lệ.')
        if box[0] < 0 or box[1] < 0 or box[2] <= box[0] or box[3] <= box[1]:
            raise ValueError('Kích thước vùng công thức không hợp lệ.')
    return images, formulas
